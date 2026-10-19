#!/usr/bin/env python3
"""Generate the ten-section Assignment 01 DOCX report from real artifacts."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPORT_PATH = Path("reports/Assignment_01_Report.docx")
METADATA_PATH = Path("models/model_metadata.json")
LEADERBOARD_PATH = Path("outputs/modeling/model_leaderboard.csv")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def add_title(document: Document, text: str, subtitle: str | None = None) -> None:
    heading = document.add_heading(text, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle:
        paragraph = document.add_paragraph(subtitle)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_page_break(document: Document) -> None:
    document.add_page_break()


def add_image_if_present(document: Document, path: Path, width: float = 5.8) -> None:
    if path.exists():
        document.add_picture(str(path), width=Inches(width))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def load_metadata() -> dict:
    if METADATA_PATH.exists():
        return json.loads(METADATA_PATH.read_text(encoding="utf-8"))
    return {
        "selected_candidate": "Run training to populate",
        "model_version": "1.0.0",
        "dataset_rows": "Pending",
        "train_rows": "Pending",
        "test_rows": "Pending",
        "metrics": {},
    }


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.color.rgb = RGBColor(23, 50, 77)
    styles["Heading 1"].font.color.rgb = RGBColor(23, 82, 122)
    styles["Heading 2"].font.color.rgb = RGBColor(41, 82, 122)
    add_page_number(section.footer.paragraphs[0])


def build_report() -> Document:
    metadata = load_metadata()
    metrics = metadata.get("metrics", {})
    document = Document()
    configure_document(document)

    document.add_heading("Machine Learning Operations (MLOps)", level=0)
    title = document.add_paragraph("Assignment 01")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.size = Pt(24)
    title.runs[0].font.bold = True
    subtitle = document.add_paragraph(
        "End-to-End Heart Disease Model Development, CI/CD, Deployment, and Monitoring"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("\nStudent: [Enter student name]\nStudent ID: [Enter student ID]\n")
    document.add_paragraph(
        "This report documents an original implementation using the UCI processed Cleveland "
        "Heart Disease dataset. The system emphasizes reproducibility, automated quality "
        "gates, portable model serving, and measurable operational behavior."
    )
    document.add_paragraph(
        f"Selected model: {metadata['selected_candidate']}\n"
        f"Model version: {metadata['model_version']}\n"
        f"Dataset records: {metadata['dataset_rows']}"
    )
    document.add_paragraph(
        "Important: this classifier is an educational demonstration and is not a medical device."
    )

    add_page_break(document)
    add_title(document, "1. Executive Summary and Architecture", "Report page 2 of 10")
    document.add_paragraph(
        "The solution converts a public clinical dataset into a versioned prediction service. "
        "A deterministic acquisition script downloads the data; cleaning converts missing "
        "markers and creates a binary target. One pipeline object contains imputation, scaling, "
        "encoding, and classification so training and serving cannot diverge."
    )
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(table.rows[0].cells, ["Stage", "Technology", "Output"]):
        cell.text = value
    architecture_rows = [
        ("Data", "Python + pandas", "Raw and cleaned CSV"),
        ("Modeling", "scikit-learn", "Complete pipeline"),
        ("Tracking", "MLflow + SQLite", "Runs, metrics, plots"),
        ("Quality", "Pytest + Ruff", "Automated gates"),
        ("Serving", "FastAPI + Docker", "REST API image"),
        ("Deployment", "Kubernetes", "Ingress, probes, HPA"),
        ("Operations", "Prometheus + Grafana", "Metrics and dashboard"),
    ]
    for row_values in architecture_rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            cell.text = value
    document.add_paragraph(
        "CI runs linting, tests, data acquisition, EDA, model training, report generation, "
        "container building, a live container smoke test, and evidence upload."
    )

    add_page_break(document)
    add_title(document, "2. Data Acquisition and Quality", "Report page 3 of 10")
    document.add_paragraph(
        "Source: UCI Machine Learning Repository, processed Cleveland Heart Disease data. "
        "The acquisition script uses an explicit HTTPS source and rejects unexpectedly small "
        "responses. The original 14 columns are assigned documented names."
    )
    add_bullets(
        document,
        [
            "Question-mark markers are converted to missing numeric values.",
            "Diagnosis values greater than zero become the positive binary class.",
            "Duplicate rows are removed and the output index is rebuilt.",
            "Age and cholesterol receive basic domain sanity checks.",
            "A JSON profile records row count, missingness, duplicates, and class balance.",
            "The cleaned CSV remains an auditable project artifact.",
        ],
    )
    document.add_paragraph(
        f"Observed artifact size: {metadata['dataset_rows']} rows. The split reserves "
        f"{metadata['test_rows']} rows for a stratified holdout and uses "
        f"{metadata['train_rows']} rows for model development."
    )
    document.add_heading("Data governance", level=2)
    document.add_paragraph(
        "The dataset is public and contains no direct identifiers in this project copy. "
        "The pipeline still treats clinical attributes as sensitive, avoids request-body logs, "
        "and reports only operational metadata."
    )

    add_page_break(document)
    add_title(document, "3. Exploratory Data Analysis", "Report page 4 of 10")
    document.add_paragraph(
        "EDA is generated by a script rather than manually edited notebook output. This makes "
        "the figures reproducible in local and CI environments. Histograms compare important "
        "continuous features by target; a correlation heatmap exposes linear relationships; "
        "and the class-balance chart checks whether accuracy alone could be misleading."
    )
    add_image_if_present(document, Path("outputs/eda/class_balance.png"), width=4.8)
    document.add_paragraph(
        "The accompanying descriptive-statistics and data-quality CSV files preserve exact "
        "values behind the charts. Missing values are handled inside the deployed model pipeline "
        "to prevent training-serving skew."
    )

    add_page_break(document)
    add_title(document, "4. Feature Engineering and Reproducibility", "Report page 5 of 10")
    document.add_paragraph(
        "Five continuous measurements (age, resting blood pressure, cholesterol, maximum heart "
        "rate, and ST depression) use median imputation followed by standard scaling. Eight "
        "categorical or ordinal-coded measurements use most-frequent imputation and one-hot "
        "encoding with unknown-category tolerance."
    )
    add_bullets(
        document,
        [
            "All transformations are fitted using training data only.",
            "The preprocessing transformer and estimator are serialized together.",
            "Feature order is stored in model metadata and enforced during inference.",
            "A fixed seed of 42 controls splitting and stochastic estimators.",
            "Pinned dependencies make clean-environment recreation repeatable.",
            "The API applies Pydantic range and type validation before inference.",
        ],
    )
    document.add_paragraph(
        "This design eliminates the common production defect in which a model receives data that "
        "was encoded or scaled differently from its training input."
    )

    add_page_break(document)
    add_title(document, "5. Model Development and Evaluation", "Report page 6 of 10")
    document.add_paragraph(
        "The experiment compares three regularized logistic-regression settings and two balanced "
        "random-forest configurations. Each candidate receives five-fold stratified cross-"
        "validation and final evaluation on an untouched 20% holdout. ROC-AUC is the selection "
        "metric, while accuracy, precision, and recall expose threshold behavior."
    )
    if LEADERBOARD_PATH.exists():
        leaderboard = pd.read_csv(LEADERBOARD_PATH)
        result_table = document.add_table(rows=1, cols=5)
        headers = ["Candidate", "CV AUC", "Accuracy", "Recall", "Test AUC"]
        for cell, value in zip(result_table.rows[0].cells, headers):
            cell.text = value
        for _, record in leaderboard.iterrows():
            cells = result_table.add_row().cells
            values = [
                record["candidate"],
                f"{record['cv_roc_auc_mean']:.3f}",
                f"{record['test_accuracy']:.3f}",
                f"{record['test_recall']:.3f}",
                f"{record['test_roc_auc']:.3f}",
            ]
            for cell, value in zip(cells, values):
                cell.text = str(value)
    document.add_paragraph(
        f"The selected candidate is {metadata['selected_candidate']}, with holdout ROC-AUC "
        f"{metrics.get('test_roc_auc', 'pending')}. Selection is automated and traceable to its "
        "MLflow run ID. Clinical use would require external validation, calibration analysis, "
        "fairness review, and medical governance beyond this coursework."
    )

    add_page_break(document)
    add_title(document, "6. Experiment Tracking and Packaging", "Report page 7 of 10")
    document.add_paragraph(
        "Every candidate executes inside an MLflow run. Parameters include model family, tuned "
        "settings, seed, split ratio, fold count, and dataset size. Metrics include cross-"
        "validation means and holdout scores. Each run logs confusion-matrix and ROC plots plus "
        "the complete scikit-learn pipeline."
    )
    add_bullets(
        document,
        [
            "Local tracking URI uses SQLite for a portable classroom setup.",
            "The leaderboard exports run IDs for review without requiring the UI.",
            "The winning joblib file contains preprocessing and prediction logic.",
            "Metadata captures feature order, model version, selection metric, and scores.",
            "The same requirements file supports training and the serving container.",
        ],
    )
    document.add_paragraph(
        "Review command: mlflow ui --backend-store-uri sqlite:///mlflow.db --port 5000. "
        "A real MLflow UI screenshot should be added to the evidence folder before submission."
    )

    add_page_break(document)
    add_title(document, "7. Automated Testing and CI/CD", "Report page 8 of 10")
    document.add_paragraph(
        "Pytest covers missing-value cleaning, target conversion, stratified splitting, complete "
        "pipeline probability output, API prediction, and invalid-request rejection. Ruff provides "
        "a fast static quality gate. The workflow fails immediately when linting, tests, training, "
        "container build, health check, or sample inference fails."
    )
    add_bullets(
        document,
        [
            "Trigger: main-branch pushes, pull requests, and manual dispatch.",
            "Training and EDA run from a clean hosted runner.",
            "The Docker image is started and queried through a real HTTP endpoint.",
            "Logs, plots, metadata, test XML, prediction output, and report are uploaded.",
            "Artifact retention is set to 30 days for reproducible grading evidence.",
        ],
    )
    add_image_if_present(document, Path("screenshots/01_ci_success.png"), width=5.8)

    add_page_break(document)
    add_title(document, "8. API, Container, and Kubernetes Deployment", "Report page 9 of 10")
    document.add_paragraph(
        "FastAPI exposes health, prediction, OpenAPI documentation, and Prometheus metrics. The "
        "prediction contract accepts all 13 model features and returns class, readable risk label, "
        "confidence, positive-class probability, and model version. The container runs as an "
        "unprivileged user and includes a health check."
    )
    document.add_paragraph(
        "Kubernetes manifests provide two replicas, readiness and liveness probes, CPU/memory "
        "boundaries, a read-only non-root security context, ClusterIP service, NGINX Ingress, and "
        "horizontal autoscaling. Local Minikube access instructions use port forwarding for a "
        "repeatable verification path and Ingress for the required exposure demonstration."
    )
    add_image_if_present(document, Path("screenshots/05_k8s_pods.png"), width=5.8)

    add_page_break(document)
    add_title(document, "9. Monitoring, Evidence, and Conclusion", "Report page 10 of 10")
    document.add_paragraph(
        "Request logs capture method, path, status, and latency without recording patient payloads. "
        "Prometheus collects request rate, response status, latency histograms, and predictions by "
        "class. Grafana is provisioned with a four-panel dashboard for operational review."
    )
    add_bullets(
        document,
        [
            "The screenshot checklist prevents missing or ambiguous execution proof.",
            "The recording guide demonstrates data, MLflow, CI, API, Kubernetes, and monitoring.",
            "A local access path is supplied when no public cloud URL is available.",
            "No screenshots, workflow success, public URL, or deployment result is fabricated.",
        ],
    )
    add_image_if_present(document, Path("screenshots/08_grafana_dashboard.png"), width=5.6)
    document.add_paragraph(
        "Conclusion: the project supplies a complete automated path from public data to an "
        "observable, containerized API. Before final submission, the student must enter identity "
        "details, publish the repository, run the target deployment, capture the listed evidence, "
        "and record the short demonstration video."
    )
    document.add_heading("References", level=2)
    document.add_paragraph(
        "UCI Machine Learning Repository: Heart Disease dataset. Documentation for MLflow, "
        "scikit-learn, FastAPI, Docker, Kubernetes, Prometheus, Grafana, and GitHub Actions."
    )
    return document


def main() -> None:
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = build_report()
    document.save(REPORT_PATH)
    print(f"Generated {REPORT_PATH}")


if __name__ == "__main__":
    main()

